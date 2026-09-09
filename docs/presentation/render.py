#!/usr/bin/env python3
"""Render the Artifacts Festival talk from `deck.py` into every target format.

    python3 docs/presentation/render.py              # all four artifacts
    python3 docs/presentation/render.py --ascii      # just slides.txt
    python3 docs/presentation/render.py --check      # validate, write nothing

Targets:

    slides.txt          78-column ASCII deck, form-feed separated
    slides.html         self-contained deck: keyboard nav, timer, notes
    slides.pptx         16:9 deck with the script in the speaker notes
    speaker_script.md   timed script with cut list and question prep

Validation runs before anything is written: body lines must fit the ASCII
frame, spoken word counts must sit within `WORD_TOLERANCE` of each slide's
time budget, and slide windows must be contiguous.
"""

from __future__ import annotations

import argparse
import html
import sys
import textwrap
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from deck import (  # noqa: E402
    QUESTIONS,
    RUNNING_LONG,
    RUNNING_SHORT,
    SLIDES,
    TEXT_WIDTH,
    TOTAL_SECONDS,
    WORD_TOLERANCE,
    WORDS_PER_MINUTE,
    Bullets,
    Para,
    Pre,
    Slide,
)

HERE = Path(__file__).resolve().parent
FRAME_INNER = TEXT_WIDTH + 4          # columns between the border pipes
FRAME_WIDTH = FRAME_INNER + 2


# --------------------------------------------------------------------------
# Shared body layout, consumed by the ASCII and PPTX renderers
# --------------------------------------------------------------------------

def body_lines(slide: Slide) -> list[str]:
    """Flatten a slide's blocks into plain monospace lines."""
    out: list[str] = []
    for index, block in enumerate(slide.blocks):
        if index:
            out.append("")
        if isinstance(block, Pre):
            out.extend(block.lines)
        elif isinstance(block, Para):
            out.extend(textwrap.wrap(block.text, width=TEXT_WIDTH) or [""])
        elif isinstance(block, Bullets):
            for item in block.items:
                wrapped = textwrap.wrap(
                    item,
                    width=TEXT_WIDTH,
                    initial_indent="  . ",
                    subsequent_indent="    ",
                )
                out.extend(wrapped)
        else:
            raise TypeError(f"unknown block type: {type(block)!r}")
    return out


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------

def validate() -> list[str]:
    problems: list[str] = []

    if TOTAL_SECONDS != 480:
        problems.append(f"deck runs {TOTAL_SECONDS}s, expected 480s (8 min)")

    previous_end = 0
    for number, slide in enumerate(SLIDES, start=1):
        if slide.start != previous_end:
            problems.append(
                f"slide {number:02d}: starts at {slide.start}s but the previous "
                f"slide ends at {previous_end}s"
            )
        previous_end = slide.end

        for line in body_lines(slide):
            if len(line) > TEXT_WIDTH:
                problems.append(
                    f"slide {number:02d}: body line is {len(line)} columns, "
                    f"max is {TEXT_WIDTH}: {line!r}"
                )

        drift = slide.spoken_words - slide.word_budget
        if abs(drift) > WORD_TOLERANCE:
            problems.append(
                f"slide {number:02d}: script is {slide.spoken_words} words but the "
                f"{slide.seconds}s window budgets {slide.word_budget} "
                f"(drift {drift:+d}, tolerance +/-{WORD_TOLERANCE})"
            )

    return problems


def report_budget() -> None:
    total = sum(s.spoken_words for s in SLIDES)
    print(f"{'#':>3}  {'window':<15} {'secs':>5} {'words':>6} {'budget':>7} {'drift':>6}")
    for number, slide in enumerate(SLIDES, start=1):
        print(
            f"{number:>3}  {slide.window:<15} {slide.seconds:>5} "
            f"{slide.spoken_words:>6} {slide.word_budget:>7} "
            f"{slide.spoken_words - slide.word_budget:>+6}"
        )
    print(
        f"\ntotal {total} spoken words -> {total / WORDS_PER_MINUTE:.2f} min at "
        f"{WORDS_PER_MINUTE} wpm, over a {TOTAL_SECONDS / 60:.0f} min slot"
    )


# --------------------------------------------------------------------------
# ASCII
# --------------------------------------------------------------------------

def render_ascii() -> str:
    pages = []
    for number, slide in enumerate(SLIDES, start=1):
        heavy = "+" + "=" * FRAME_INNER + "+"
        thin = "+" + "-" * FRAME_INNER + "+"
        left = f"{number:02d} / {len(SLIDES)}"
        pad = TEXT_WIDTH - len(left) - len(slide.window)
        rows = [
            heavy,
            f"|  {left}{' ' * pad}{slide.window}  |",
            f"|  {slide.title.ljust(TEXT_WIDTH)}  |",
            thin,
            f"|  {' ' * TEXT_WIDTH}  |",
        ]
        for line in body_lines(slide):
            rows.append(f"|  {line.ljust(TEXT_WIDTH)}  |")
        rows.append(f"|  {' ' * TEXT_WIDTH}  |")
        rows.append(heavy)
        pages.append("\n".join(rows) + "\n")
    return "\f\n".join(pages)


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------

HTML_CSS = """
:root {
  --bg: #0d1117;
  --panel: #12181f;
  --ink: #d5dde5;
  --dim: #7d8b99;
  --accent: #5ec8f2;
  --rule: #263139;
  --warn: #e8b046;
  --late: #e0685f;
  --mono: ui-monospace, "SF Mono", "Cascadia Mono", Menlo, Consolas,
          "DejaVu Sans Mono", "Liberation Mono", monospace;
}
* { box-sizing: border-box; }
html, body {
  margin: 0; height: 100%; background: var(--bg); color: var(--ink);
  font-family: var(--mono); overflow: hidden;
}
#deck { height: 100%; }
.slide {
  display: none; height: 100%; padding: 3.2vh 3.5vw 6vh;
  flex-direction: column;
  /* Fit the widest line (72 chars) and the tallest slide at once. */
  font-size: min(calc(93vw / 46), calc(88vh / var(--lines)));
}
.slide.active { display: flex; }
.slide-head {
  display: flex; justify-content: space-between; align-items: baseline;
  gap: 2em; border-bottom: 1px solid var(--rule);
  padding-bottom: 0.5em; margin-bottom: 1.2em; flex: none;
}
.slide-title {
  font-size: 1.5em; font-weight: 700; letter-spacing: 0.04em;
  color: var(--accent); text-transform: uppercase;
}
.slide-meta { font-size: 0.95em; color: var(--dim); white-space: nowrap; }
.slide-body { flex: 1 1 auto; min-height: 0; }
.slide-body > * { margin: 0 0 1.15em; }
.slide-body > *:last-child { margin-bottom: 0; }
pre {
  font-family: var(--mono); font-size: 1em; line-height: 1.28;
  margin: 0 0 1.15em; white-space: pre; overflow: visible;
  color: var(--ink);
}
/* Match the 72-column authoring width so prose breaks like the ASCII deck. */
p { line-height: 1.5; max-width: 74ch; }
ul { margin: 0 0 1.15em; padding-left: 1.6em; }
li { line-height: 1.5; margin-bottom: 0.35em; max-width: 70ch; }

/* Title slide */
.slide.title { justify-content: center; align-items: center; text-align: center; }
.slide.title .slide-head { display: none; }
.slide.title .slide-body > * { margin-bottom: 1.4em; }
.slide.title pre { color: var(--accent); }
.slide.title p { max-width: none; }

/* Heads-up display */
#hud {
  position: fixed; left: 0; right: 0; bottom: 0; display: flex;
  justify-content: space-between; align-items: center; gap: 1.5em;
  padding: 0.45em 1.2em; font-size: 0.78rem; color: var(--dim);
  background: var(--panel); border-top: 1px solid var(--rule);
}
#hud b { color: var(--ink); font-weight: 600; }
#clock.behind { color: var(--warn); }
#clock.late { color: var(--late); }
#bar {
  position: fixed; left: 0; bottom: 0; height: 2px;
  background: var(--accent); transition: width 0.18s ease;
}

/* Speaker notes */
#notes {
  position: fixed; right: 0; top: 0; bottom: 2.1em; width: min(34rem, 42vw);
  padding: 1.4em 1.6em; overflow-y: auto; font-size: 0.82rem; line-height: 1.55;
  background: var(--panel); border-left: 1px solid var(--rule);
  transform: translateX(101%); transition: transform 0.2s ease;
}
#notes.open { transform: translateX(0); }
#notes h3 {
  margin: 0 0 0.9em; font-size: 0.8rem; letter-spacing: 0.1em;
  text-transform: uppercase; color: var(--accent);
}
#notes p { max-width: none; margin: 0 0 0.9em; }
#notes .cue {
  color: var(--warn); border-left: 2px solid var(--warn);
  padding-left: 0.8em; margin-bottom: 0.9em;
}
#notes .opt {
  color: var(--dim); border-left: 2px solid var(--rule);
  padding-left: 0.8em;
}

@media print {
  html, body { overflow: visible; height: auto; background: #fff; color: #111; }
  #hud, #notes, #bar { display: none; }
  .slide {
    display: flex !important; height: 100vh; page-break-after: always;
    font-size: 11pt;
  }
  .slide-title { color: #111; }
  .slide.title pre { color: #111; }
}
"""

HTML_JS = """
const slides = Array.from(document.querySelectorAll('.slide'));
const notes = document.getElementById('notes');
const bar = document.getElementById('bar');
const counter = document.getElementById('counter');
const target = document.getElementById('target');
const clock = document.getElementById('clock');
const TOTAL = Number(document.body.dataset.total);
let index = 0;
let started = null;

function show(next) {
  index = Math.max(0, Math.min(slides.length - 1, next));
  slides.forEach((s, i) => s.classList.toggle('active', i === index));
  const slide = slides[index];
  counter.textContent = `${String(index + 1).padStart(2, '0')} / ${slides.length}`;
  target.textContent = slide.dataset.window;
  bar.style.width = `${((index + 1) / slides.length) * 100}%`;
  notes.innerHTML = slide.querySelector('template').innerHTML;
  location.hash = index + 1;
}

function tick() {
  if (started === null) { clock.textContent = '--:-- idle'; return; }
  const elapsed = Math.floor((Date.now() - started) / 1000);
  const mm = String(Math.floor(elapsed / 60)).padStart(2, '0');
  const ss = String(elapsed % 60).padStart(2, '0');
  clock.textContent = `${mm}:${ss}`;
  // Compare against where this slide should have ended.
  const due = Number(slides[index].dataset.end);
  clock.className = elapsed > TOTAL ? 'late' : (elapsed > due ? 'behind' : '');
}

document.addEventListener('keydown', (event) => {
  const key = event.key;
  if (['ArrowRight', 'ArrowDown', 'PageDown', ' ', 'Enter'].includes(key)) {
    if (started === null) started = Date.now();
    show(index + 1); event.preventDefault();
  } else if (['ArrowLeft', 'ArrowUp', 'PageUp'].includes(key)) {
    show(index - 1); event.preventDefault();
  } else if (key === 'Home') { show(0); }
  else if (key === 'End') { show(slides.length - 1); }
  else if (key === 's' || key === 'S') { notes.classList.toggle('open'); }
  else if (key === 'f' || key === 'F') {
    if (document.fullscreenElement) document.exitFullscreen();
    else document.documentElement.requestFullscreen();
  } else if (key === 't' || key === 'T') {
    started = started === null ? Date.now() : null;
  } else if (key === 'r' || key === 'R') { started = Date.now(); }
});

document.addEventListener('click', (event) => {
  if (event.target.closest('#notes') || event.target.closest('#hud')) return;
  if (started === null) started = Date.now();
  show(index + 1);
});

show(Number(location.hash.slice(1)) - 1 || 0);
setInterval(tick, 250);
tick();
"""


def _html_blocks(slide: Slide) -> str:
    parts = []
    for block in slide.blocks:
        if isinstance(block, Pre):
            parts.append("<pre>" + html.escape("\n".join(block.lines)) + "</pre>")
        elif isinstance(block, Para):
            parts.append("<p>" + html.escape(block.text) + "</p>")
        elif isinstance(block, Bullets):
            items = "".join(f"<li>{html.escape(i)}</li>" for i in block.items)
            parts.append(f"<ul>{items}</ul>")
    return "\n".join(parts)


def _html_notes(slide: Slide, number: int) -> str:
    parts = [f"<h3>Slide {number:02d} &middot; {html.escape(slide.window)}</h3>"]
    if slide.delivery:
        parts.append(f'<p class="cue">{html.escape(slide.delivery)}</p>')
    parts.extend(f"<p>{html.escape(p)}</p>" for p in slide.script)
    if slide.optional:
        parts.append(
            f'<p class="opt">Only if ahead of the clock: '
            f'&ldquo;{html.escape(slide.optional)}&rdquo;</p>'
        )
    parts.append(
        f"<p class=\"opt\">{slide.spoken_words} words for a {slide.seconds}s "
        f"window (budget {slide.word_budget}).</p>"
    )
    return "\n".join(parts)


def render_html() -> str:
    sections = []
    for number, slide in enumerate(SLIDES, start=1):
        # Drive the auto-fit font size off this slide's own height.
        lines = round((len(body_lines(slide)) + 7) * 1.3)
        classes = "slide title" if slide.kind == "title" else "slide"
        sections.append(
            f'<section class="{classes}" style="--lines: {lines}"'
            f' data-window="{html.escape(slide.window)}" data-end="{slide.end}">\n'
            f'  <div class="slide-head">\n'
            f'    <div class="slide-title">{html.escape(slide.title)}</div>\n'
            f'    <div class="slide-meta">{number:02d} / {len(SLIDES)}'
            f' &middot; {html.escape(slide.window)}</div>\n'
            f'  </div>\n'
            f'  <div class="slide-body">\n{_html_blocks(slide)}\n  </div>\n'
            f'  <template>{_html_notes(slide, number)}</template>\n'
            f'</section>'
        )

    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Grading Is Easy, Feedback Is Hard &mdash; Artifacts Festival</title>
<style>{HTML_CSS}</style>
</head>
<body data-total="{TOTAL_SECONDS}">
<div id="deck">
{chr(10).join(sections)}
</div>
<div id="bar"></div>
<div id="hud">
  <span><b id="counter">01 / {len(SLIDES)}</b> &nbsp; target <b id="target"></b></span>
  <span>arrows/space move &middot; <b>s</b> notes &middot; <b>t</b> timer &middot;
        <b>r</b> reset &middot; <b>f</b> fullscreen</span>
  <span>elapsed <b id="clock">--:--</b> / {_mmss(TOTAL_SECONDS)}</span>
</div>
<aside id="notes"></aside>
<script>{HTML_JS}</script>
</body>
</html>
"""


def _mmss(seconds: int) -> str:
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------

BG = 0x0D1117
INK = 0xD5DDE5
ACCENT = 0x5EC8F2
DIM = 0x7D8B99
MONO = "Courier New"        # chosen for availability: ASCII tables must not reflow


def render_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(value: int) -> RGBColor:
        return RGBColor((value >> 16) & 0xFF, (value >> 8) & 0xFF, value & 0xFF)

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    body_left, body_top = Inches(0.6), Inches(1.32)
    body_w, body_h = Inches(12.13), Inches(5.75)
    line_spacing = 1.16

    for number, slide_spec in enumerate(SLIDES, start=1):
        slide = deck.slides.add_slide(blank)

        background = slide.background
        background.fill.solid()
        background.fill.fore_color.rgb = rgb(BG)

        lines = body_lines(slide_spec)
        is_title = slide_spec.kind == "title"

        if not is_title:
            head = slide.shapes.add_textbox(
                body_left, Inches(0.42), body_w, Inches(0.85)
            ).text_frame
            head.word_wrap = True
            run = head.paragraphs[0].add_run()
            run.text = slide_spec.title
            run.font.name = MONO
            run.font.size = Pt(26)
            run.font.bold = True
            run.font.color.rgb = rgb(ACCENT)

            meta = slide.shapes.add_textbox(
                body_left, Inches(1.02), body_w, Inches(0.32)
            ).text_frame
            meta_run = meta.paragraphs[0].add_run()
            meta_run.text = f"{number:02d} / {len(SLIDES)}   {slide_spec.window}"
            meta_run.font.name = MONO
            meta_run.font.size = Pt(12)
            meta_run.font.color.rgb = rgb(DIM)

        top = Inches(0.95) if is_title else body_top
        height = Inches(5.9) if is_title else body_h

        # Auto-fit: Courier New advances 0.6em per character.
        widest = max((len(l) for l in lines), default=1) or 1
        avail_w_pt = body_w / Emu(1) / 914400 * 72
        avail_h_pt = height / Emu(1) / 914400 * 72
        by_width = avail_w_pt / (widest * 0.60)
        by_height = avail_h_pt / (max(len(lines), 1) * line_spacing)
        size = max(9.0, min(by_width, by_height, 22.0 if is_title else 16.0))

        box = slide.shapes.add_textbox(body_left, top, body_w, height)
        frame = box.text_frame
        frame.word_wrap = False

        for position, line in enumerate(lines):
            para = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
            para.line_spacing = line_spacing
            if is_title:
                para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            run.text = line
            run.font.name = MONO
            run.font.size = Pt(size)
            run.font.color.rgb = rgb(ACCENT if is_title and position < 9 else INK)

        notes = slide.notes_slide.notes_text_frame
        chunks = [f"[{slide_spec.window}]  slide {number:02d} of {len(SLIDES)}"]
        if slide_spec.delivery:
            chunks.append(f"CUE: {slide_spec.delivery}")
        chunks.extend(slide_spec.script)
        if slide_spec.optional:
            chunks.append(f"IF AHEAD: {slide_spec.optional}")
        chunks.append(
            f"({slide_spec.spoken_words} words / {slide_spec.seconds}s window, "
            f"budget {slide_spec.word_budget})"
        )
        notes.text = "\n\n".join(chunks)

    deck.save(path)


# --------------------------------------------------------------------------
# Speaker script
# --------------------------------------------------------------------------

def render_script() -> str:
    total = sum(s.spoken_words for s in SLIDES)
    out = [
        "# Artifacts Festival — 8-minute speaker script",
        "",
        "<!-- Generated by render.py from deck.py. Edit deck.py, not this file. -->",
        "",
        "Companion to [`slides.txt`](slides.txt), [`slides.html`](slides.html) and",
        "[`slides.pptx`](slides.pptx) — the PPTX carries this script in its speaker",
        f"notes. Target pace is {WORDS_PER_MINUTE} words per minute; the spoken text is",
        f"{total} words over a {TOTAL_SECONDS // 60}-minute slot.",
        "",
        "| Slide | Window | Section | Words |",
        "| ---: | :--- | :--- | ---: |",
    ]
    for number, slide in enumerate(SLIDES, start=1):
        out.append(
            f"| {number:02d} | {slide.window} | {slide.title.title()} "
            f"| {slide.spoken_words} |"
        )
    out.append("")

    for number, slide in enumerate(SLIDES, start=1):
        out += ["---", "", f"## Slide {number:02d} — {slide.title.title()} · {slide.window}", ""]
        if slide.delivery:
            out += [f"> Delivery: {slide.delivery}", ""]
        for paragraph in slide.script:
            out += textwrap.wrap(paragraph, width=79) + [""]
        if slide.optional:
            out += textwrap.wrap(
                f'> Optional, only if ahead of the clock: "{slide.optional}"',
                width=79,
                subsequent_indent="> ",
            ) + [""]

    out += ["---", "", "## If you are running long", "",
            "Cut in this order. Each cut is self-contained and costs no continuity.", ""]
    for position, item in enumerate(RUNNING_LONG, start=1):
        out += textwrap.wrap(
            f"{position}. {item}", width=79, subsequent_indent="   "
        )
    out += [
        "",
        "Do not cut slides 04, 05 or 10. The prompt structure, the injection "
        "mechanism and the one-command reproduction are what the audience came for.",
        "",
        "## If you are running short",
        "",
    ]
    for item in RUNNING_SHORT:
        out += textwrap.wrap(item, width=79) + [""]

    out += ["## Question prep", ""]
    for question, answer in QUESTIONS:
        out += [f"**{question}**"]
        out += textwrap.wrap(answer, width=79) + [""]

    return "\n".join(out).rstrip("\n") + "\n"


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--ascii", action="store_true", help="write slides.txt")
    parser.add_argument("--html", action="store_true", help="write slides.html")
    parser.add_argument("--pptx", action="store_true", help="write slides.pptx")
    parser.add_argument("--script", action="store_true", help="write speaker_script.md")
    parser.add_argument("--check", action="store_true", help="validate only")
    parser.add_argument("--budget", action="store_true", help="print the word budget")
    args = parser.parse_args()

    problems = validate()
    if problems:
        print("deck validation failed:", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        return 1

    if args.budget:
        report_budget()
    if args.check:
        print(f"deck ok: {len(SLIDES)} slides, {TOTAL_SECONDS}s")
        return 0

    selected = args.ascii or args.html or args.pptx or args.script
    written = []

    if args.ascii or not selected:
        path = HERE / "slides.txt"
        path.write_text(render_ascii(), encoding="ascii")
        written.append(path)
    if args.html or not selected:
        path = HERE / "slides.html"
        path.write_text(render_html(), encoding="utf-8")
        written.append(path)
    if args.pptx or not selected:
        path = HERE / "slides.pptx"
        try:
            render_pptx(path)
        except ImportError:
            print(
                "python-pptx is not installed; skipping slides.pptx\n"
                "  python3 -m pip install -r docs/presentation/requirements.txt",
                file=sys.stderr,
            )
        else:
            written.append(path)
    if args.script or not selected:
        path = HERE / "speaker_script.md"
        path.write_text(render_script(), encoding="utf-8")
        written.append(path)

    for path in written:
        print(f"wrote {path.relative_to(HERE.parents[1])} ({path.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
